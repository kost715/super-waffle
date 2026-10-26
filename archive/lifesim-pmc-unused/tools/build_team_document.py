#!/usr/bin/env python3
"""Build the team onboarding document as PDF and editable DOCX from one Markdown.

Requires pandoc, reportlab, python-docx and Noto fonts. No game files are changed.
"""
from __future__ import annotations
import hashlib
import html
import json
from pathlib import Path
import re
import subprocess
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfdoc, pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import BaseDocTemplate, CondPageBreak, Flowable, Frame, NextPageTemplate, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.platypus.tableofcontents import TableOfContents

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/TEAM_DEVELOPMENT_PLAN_RU.md'
DIST = ROOT / 'dist'
REPORT = ROOT / 'reports/team-document'
PDF = DIST / 'LifeSim_PMC_Team_Development_Plan_RU.pdf'
DOCX = DIST / 'LifeSim_PMC_Team_Development_Plan_RU.docx'
W, H = A4
M = 46
CW = W - 2*M
NAVY = colors.HexColor('#142E3A')
TEAL = colors.HexColor('#087D6B')
MINT = colors.HexColor('#4DD9B2')
INK = colors.HexColor('#253B44')
MUTED = colors.HexColor('#5D717B')
LIGHT = colors.HexColor('#F1F6F6')
RULE = colors.HexColor('#D8E4E6')


def fonts():
    folder = Path('/usr/share/fonts/truetype/noto')
    for name, file in [('Body','NotoSans-Regular.ttf'), ('Bold','NotoSans-Bold.ttf'), ('Italic','NotoSans-Italic.ttf'), ('BoldItalic','NotoSans-BoldItalic.ttf'), ('Display','NotoSansDisplay-Bold.ttf')]:
        pdfmetrics.registerFont(TTFont(name, str(folder/file)))
    pdfmetrics.registerFont(TTFont('Symbols','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
    pdfmetrics.registerFontFamily('Body', normal='Body', bold='Bold', italic='Italic', boldItalic='BoldItalic')


fonts()
STYLES = {
    'body': ParagraphStyle('Body',fontName='Body',fontSize=9.6,leading=14.3,textColor=INK,spaceAfter=7,allowWidows=0,allowOrphans=0),
    'h1': ParagraphStyle('Title',fontName='Display',fontSize=25,leading=32,textColor=NAVY,spaceAfter=17,keepWithNext=True),
    'h2': ParagraphStyle('Chapter',fontName='Display',fontSize=16.5,leading=22,textColor=NAVY,spaceBefore=16,spaceAfter=10,keepWithNext=True),
    'h3': ParagraphStyle('Section',fontName='Bold',fontSize=11.4,leading=16.4,textColor=TEAL,spaceBefore=10,spaceAfter=7,keepWithNext=True),
    'cell': ParagraphStyle('Cell',fontName='Body',fontSize=8.6,leading=12.6,textColor=INK),
    'cellhead': ParagraphStyle('CellHead',fontName='Bold',fontSize=8.7,leading=12.8,textColor=colors.white),
    'small': ParagraphStyle('Small',fontName='Body',fontSize=8.7,leading=13,textColor=MUTED,spaceAfter=8),
}


def inline(nodes, formatted=True):
    result=[]
    for node in nodes:
        kind=node['t']; val=node.get('c')
        if kind=='Str': result.append(html.escape(val).replace('→','<font name="Symbols">→</font>') if formatted else val)
        elif kind in ['Space','SoftBreak']: result.append(' ')
        elif kind=='LineBreak': result.append('<br/>' if formatted else ' ')
        elif kind in ['Strong','Emph']:
            text=inline(val,formatted); tag='b' if kind=='Strong' else 'i'
            result.append(f'<{tag}>{text}</{tag}>' if formatted else text)
        elif kind=='Code': result.append(html.escape(val[1]) if formatted else val[1])
        elif kind=='Link': result.append(inline(val[1],formatted))
        elif kind=='Quoted': result.append('«'+inline(val[1],formatted)+'»')
        else: raise ValueError(f'Unsupported inline: {kind}')
    return ''.join(result)


def para(text, style='body', **kwargs):
    chosen=STYLES[style]
    if kwargs: chosen=ParagraphStyle(chosen.name+'Custom',parent=chosen,**kwargs)
    return Paragraph(text,chosen)


def draw(c,text,x,top,width,size=10,font='Body',color=INK,leading=None):
    p=Paragraph(text,ParagraphStyle('Canvas',fontName=font,fontSize=size,leading=leading or size*1.45,textColor=color))
    _,height=p.wrap(width,H);p.drawOn(c,x,top-height);return height


class Cover(Flowable):
    def __init__(self):
        super().__init__()
        self.width=W;self.height=H
    def draw(self):
        c=self.canv
        c.bookmarkPage('cover');c.addOutlineEntry('LifeSim PMC — план команды','cover',0)
        c.setFillColor(NAVY);c.rect(0,0,W,H,stroke=0,fill=1)
        c.setFillColor(MINT);c.rect(M,H-69,44,4,stroke=0,fill=1)
        draw(c,'ARMA LIFESIM / PMC',M,H-36,CW,size=10,font='Bold',color=colors.white)
        draw(c,'ПЛАН ДЛЯ КОМАНДЫ',M,697,CW,size=10,font='Bold',color=MINT)
        draw(c,'Этапы<br/>разработки',M,650,CW,size=42,font='Display',color=colors.white,leading=52)
        draw(c,'История проекта, текущая сборка<br/>и порядок дальнейшей работы',M,508,CW,size=17,color=colors.HexColor('#CCE3E3'),leading=26)
        draw(c,'LifeSim PMC — этапы разработки и план работы команды',M,409,CW,size=10.2,font='Bold',color=colors.white)
        draw(c,'Для разработчиков SQF, интерфейсов,<br/>геймдизайнеров и тестировщиков.',M,377,CW,size=11,color=colors.HexColor('#CCE3E3'),leading=18)
        cards=[('0.6.1','Опорная alpha-сборка'),('10','Типов основных контрактов'),('M0–M8','Следующие этапы')]
        gap=11;cw=(CW-2*gap)/3
        for i,(n,label) in enumerate(cards):
            x=M+i*(cw+gap);c.setFillColor(colors.HexColor('#244651'));c.roundRect(x,191,cw,101,7,stroke=0,fill=1)
            draw(c,n,x+12,279,cw-24,size=27,font='Display',color=MINT)
            draw(c,label,x+12,233,cw-24,size=8.7,color=colors.white)
        draw(c,'Реализованное и запланированное разделены.<br/>План не задаёт неподтверждённые сроки выпуска.',M,163,CW,size=10,color=colors.HexColor('#CCE3E3'),leading=16)
        c.setStrokeColor(colors.HexColor('#46646D'));c.line(M,84,W-M,84)
        draw(c,'Редакция 1.0 · 21 октября 2026',M,65,CW,size=9,color=colors.white)


def footer(c,doc):
    c.saveState();c.setStrokeColor(RULE);c.setLineWidth(.5)
    c.line(M,H-44,W-M,H-44);c.line(M,37,W-M,37)
    c.setFont('Bold',7.4);c.setFillColor(NAVY);c.drawString(M,H-30,'LIFESIM PMC / ПЛАН РАЗРАБОТКИ')
    c.setFont('Body',7.3);c.setFillColor(MUTED);c.drawRightString(W-M,H-30,'ДЛЯ КОМАНДЫ · РЕДАКЦИЯ 1.0')
    c.drawString(M,23,'Опорная сборка 0.6.1 alpha 1 · 21.10.2026')
    c.setFont('Bold',8);c.drawRightString(W-M,23,f'{doc.page:02d}');c.restoreState()


class PdfCanvas(canvas.Canvas):
    def __init__(self,*a,**kw):
        super().__init__(*a,**kw)
        self.setTitle('LifeSim PMC — этапы разработки и план работы команды')
        self.setAuthor('LifeSim PMC');self.setSubject('История разработки, дорожная карта, роли и критерии приёмки')
        self._doc.Catalog.Lang=pdfdoc.PDFString('ru-RU')


class TeamDoc(BaseDocTemplate):
    def afterFlowable(self,f):
        if hasattr(f,'bookmark'):
            self.canv.bookmarkPage(f.bookmark)
            self.canv.addOutlineEntry(f.getPlainText(),f.bookmark,level=f.level,closed=False)
            if f.level==1 and f.bookmark!='contents':self.notify('TOCEntry',(0,f.getPlainText(),self.page,f.bookmark))


def heading(text,style,key,level):
    f=para(text,style);f.bookmark=key;f.level=level;return f


def table_rows(value):
    rows=list(value[3][1])
    for body in value[4]:rows.extend(body[2]);rows.extend(body[3])
    rows.extend(value[5][1]);return rows


def widths(rows):
    n=len(rows[0][1]);title=inline(rows[0][1][0][4][0]['c'],False)
    if n==3:
        return [CW*.21,CW*.40,CW*.39] if 'приоритет' not in title else [CW*.19,CW*.43,CW*.38]
    if title=='Файл или каталог':return [CW*.40,CW*.60]
    if title=='До этапа':return [CW*.18,CW*.82]
    return [CW*.30,CW*.70]


def blocks_to_pdf(blocks,indent=0):
    out=[]
    for index,b in enumerate(blocks):
        kind=b['t'];v=b.get('c')
        if kind in ['Para','Plain']:
            p=para(inline(v),leftIndent=indent)
            if index+1<len(blocks) and blocks[index+1]['t'] in ['BulletList','OrderedList','Table'] and p.getPlainText().rstrip().endswith(':'):p.keepWithNext=True
            out.append(p)
        elif kind=='Header':
            level,attrs,nodes=v
            if level==1:continue
            out.append(CondPageBreak(150 if level==2 else 100))
            out.append(heading(inline(nodes),'h2' if level==2 else 'h3',attrs[0],level-1))
        elif kind in ['BulletList','OrderedList']:
            items=v if kind=='BulletList' else v[1];start=1 if kind=='BulletList' else v[0][0]
            for i,item in enumerate(items):
                first,*rest=item
                if first['t'] not in ['Para','Plain']:raise ValueError('Unexpected list item')
                sty=ParagraphStyle('List',parent=STYLES['body'],leftIndent=indent+14,bulletIndent=indent,bulletFontName='Body',bulletFontSize=9,spaceAfter=4)
                out.append(Paragraph(inline(first['c']),sty,bulletText='•' if kind=='BulletList' else f'{start+i}.'))
                out.extend(blocks_to_pdf(rest,indent+14))
            out.append(Spacer(1,4))
        elif kind=='Table':
            rows=table_rows(v)
            data=[]
            for i,row in enumerate(rows):
                data.append([para('<br/>'.join(inline(x['c']) for x in cell[4] if x['t'] in ['Para','Plain']),'cellhead' if i==0 else 'cell') for cell in row[1]])
            t=Table(data,colWidths=widths(rows),repeatRows=1,hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),NAVY),('ROWBACKGROUNDS',(0,1),(-1,-1),[LIGHT,colors.white]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),8),('RIGHTPADDING',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),('LINEBELOW',(0,-1),(-1,-1),.5,RULE)]));t.spaceAfter=12;out.append(t)
        else:raise ValueError(f'Unsupported block: {kind}')
    return out


def source_segments(blocks):
    out=[]
    for b in blocks:
        k=b['t'];v=b.get('c')
        if k in ['Para','Plain']:out.append(inline(v,False))
        elif k=='Header':out.append(inline(v[2],False))
        elif k in ['BulletList','OrderedList']:
            for item in v if k=='BulletList' else v[1]:out.extend(source_segments(item))
        elif k=='Table':
            for row in table_rows(v):
                for cell in row[1]:out.extend(source_segments(cell[4]))
        else:raise ValueError(k)
    return out


def make_pdf(ast):
    toc=TableOfContents();toc.dotsMinLevel=0
    toc.levelStyles=[ParagraphStyle('Toc',fontName='Body',fontSize=10,leading=15,textColor=INK,spaceBefore=7,rightIndent=26)]
    story=[Cover(),NextPageTemplate('body'),PageBreak(),heading('Содержание','h1','contents',1),para('Документ включает пройденные этапы, действующую сборку и согласованное развитие. Пункты содержания и закладки PDF ведут к нужным разделам.','small'),toc,Spacer(1,17),para('Первый шаг команды — подтвердить опорную 0.6.1 в игре и согласовать общую модель компании. Казна, закупки, персонал, сабконтракты и регионы пока относятся к будущей разработке.','small'),PageBreak()]
    story.extend(blocks_to_pdf(ast['blocks']))
    doc=TeamDoc(str(PDF),pagesize=A4,leftMargin=M,rightMargin=M,topMargin=62,bottomMargin=52,pageCompression=1)
    doc.addPageTemplates([PageTemplate('cover',[Frame(0,0,W,H,leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)]),PageTemplate('body',[Frame(M,52,CW,H-114,leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)],onPage=footer)])
    doc.multiBuild(story,canvasmaker=PdfCanvas)


def field(paragraph,instruction):
    run=paragraph.add_run();start=OxmlElement('w:fldChar');start.set(qn('w:fldCharType'),'begin');run._r.append(start)
    run=paragraph.add_run();text=OxmlElement('w:instrText');text.text=instruction;run._r.append(text)
    run=paragraph.add_run();end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');run._r.append(end)


def make_docx():
    ref=Document();sec=ref.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
    for name,size,bold,color in [('Normal',10.2,False,'253B44'),('Heading 1',23,True,'142E3A'),('Heading 2',16,True,'142E3A'),('Heading 3',11.5,True,'087D6B')]:
        s=ref.styles[name];s.font.name='Noto Sans';s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor.from_string(color)
        s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.15
        if name!='Normal':s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(11)
    ref.save(REPORT/'reference.docx')
    subprocess.run(['pandoc','-f','gfm',str(SOURCE),'--reference-doc',str(REPORT/'reference.docx'),'-o',str(DOCX)],check=True)
    d=Document(DOCX);d.core_properties.title='LifeSim PMC — этапы разработки и план работы команды';d.core_properties.author='LifeSim PMC';d.core_properties.subject='Дорожная карта и подключение команды';d.core_properties.language='ru-RU'
    header=d.sections[0].header.paragraphs[0];header.text='LIFESIM PMC  /  ПЛАН РАЗРАБОТКИ';header.style='Caption'
    footer=d.sections[0].footer.paragraphs[0];footer.text='Редакция 1.0 · 21.10.2026  |  Страница ';field(footer,' PAGE ')
    for t in d.tables:
        t.style='Table Grid';t.autofit=False
        for i,row in enumerate(t.rows):
            pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            if i==0:pr.append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'142E3A');cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3);p.paragraph_format.keep_with_next=(i==0)
                    for run in p.runs:
                        run.font.name='Noto Sans';run.font.size=Pt(9)
                        if i==0:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    d.save(DOCX)


def main():
    DIST.mkdir(exist_ok=True);REPORT.mkdir(parents=True,exist_ok=True)
    ast=json.loads(subprocess.run(['pandoc','-f','gfm','-t','json',str(SOURCE)],check=True,capture_output=True,text=True).stdout)
    make_pdf(ast);make_docx()
    manifest={'source':str(SOURCE.relative_to(ROOT)),'source_segments':source_segments(ast['blocks']),'outputs':{p.name:{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [PDF,DOCX]},'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest()}
    (REPORT/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in manifest.items() if k!='source_segments'},ensure_ascii=False,indent=2))


if __name__=='__main__':main()
